"""
DEMO: dien thong tin tu file Template G-xx-xxxx.xlsx (cac o vang o Sheet1) vao ECUS, tab "Thong tin chung".

Chi dung thu vien co san cua Python + openpyxl (khong can pywinauto / pywin32).

Cach dung:
1. Mo ECUS (nen thu tren ban Demo), vao Tờ khai xuất khẩu, tab "Thông tin chung",
   o che do nhap (to khai moi / dang sua, cac o khong bi khoa).
2. Chay: python ecus_fill_demo.py
3. Chon file Template trong D:\\test\\output.
4. Script se IN RA bang "gia tri se dien" va "o tim thay trong ECUS" de ban kiem tra truoc (chua dien gi).
5. Chi khi ban go DIEN thi moi tu go vao ECUS. Script KHONG bam Ghi / Gui.

Cac o con lai (booking, ngay gio...) ban tu nhap nhu cu.
"""
import ctypes
import re
import sys
import time
import unicodedata
from collections import namedtuple
from ctypes import wintypes
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

# ===== CAU HINH =====
BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "output"
SHEET = "Sheet1"

# O vang trong file Template (Sheet1)
CELLS = {
    "invoice_no": "B21",     # So invoice
    "invoice_date": "B22",   # Ngay invoice (= Ship Date)
    "currency": "C12",       # Currency
    "terms": "C11",          # Terms of sale (vi du "FOB Ho Chi Minh City" -> lay "FOB")
    "invoice_total": "B27",  # Tong tri gia hoa don
    "qty": "D16",            # So luong kien
    "unit": "D17",           # Don vi (PP, BX...)
    "gross": "D18",          # KGS
}


# ===== DINH DANG GIA TRI CHO ECUS =====
def fmt_text(v):
    return str(v).strip()


def fmt_date(v):
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def fmt_terms(v):
    parts = str(v).strip().split()
    return parts[0].upper() if parts else ""


def fmt_money(v):
    return f"{float(v):,.2f}"


def fmt_int(v):
    return str(int(round(float(v))))


def fmt_num(v):
    return ("%.3f" % float(v)).rstrip("0").rstrip(".")


# (ten hien thi, khoa o Template, nhan trong ECUS, chi so o ben phai nhan, ham dinh dang)
FIELDS = [
    ("So hoa don",           "invoice_no",    "Số hóa đơn:",                    0, fmt_text),
    ("Ngay phat hanh",       "invoice_date",  "Ngày phát hành:",                0, fmt_date),
    ("Dieu kien gia",        "terms",         "Điều kiện giá hóa đơn:",         0, fmt_terms),
    ("Ma dong tien",         "currency",      "Mã đồng tiền của hóa đơn:",      0, fmt_text),
    ("Tong tri gia hoa don", "invoice_total", "Tổng trị giá hóa đơn:",          0, fmt_money),
    ("So luong kien",        "qty",           "Số lượng kiện:",                 0, fmt_int),
    ("Don vi kien",          "unit",          "Số lượng kiện:",                 1, fmt_text),
    ("Gross (KGS)",          "gross",         "Tổng trọng lượng hàng (Gross):", 0, fmt_num),
]

Ctl = namedtuple("Ctl", "hwnd cls text left top right bottom visible")


# ===== DOC FILE TEMPLATE =====
def read_template(path):
    wb = load_workbook(path, data_only=True)
    ws = None
    for name in wb.sheetnames:
        if name.strip().casefold() == SHEET.casefold():
            ws = wb[name]
            break
    if ws is None:
        raise RuntimeError(f"File khong co sheet '{SHEET}'")
    values = {key: ws[addr].value for key, addr in CELLS.items()}
    wb.close()
    return values


# ===== TIM O TRONG ECUS =====
def norm_text(s):
    s = unicodedata.normalize("NFC", str(s or ""))
    return re.sub(r"\s+", " ", s).strip().rstrip(":").strip().casefold()


def is_edit(ctl):
    return ".EDIT." in ctl.cls.upper() or ctl.cls == "Edit"


def locate(items, label_text, index):
    """O EDIT thu 'index' (tinh tu trai sang phai) nam ben phai nhan, cung dong. Tra ve (ctl, ghi chu)."""
    target = norm_text(label_text)
    labels = [c for c in items
              if c.visible and "STATIC" in c.cls.upper() and norm_text(c.text) == target]
    if not labels:
        return None, "khong thay nhan (co dang o dung tab 'Thong tin chung' khong?)"
    label = sorted(labels, key=lambda c: (c.top, c.left))[0]
    note = f"co {len(labels)} nhan trung ten, lay cai tren cung" if len(labels) > 1 else ""

    edits = [c for c in items
             if c.visible and is_edit(c)
             and c.left >= label.right - 2
             and c.top <= label.bottom and c.bottom >= label.top]
    edits.sort(key=lambda c: c.left)
    if len(edits) <= index:
        return None, "khong thay o nhap ben phai nhan"
    return edits[index], note


def same_value(a, b):
    a, b = str(a).strip(), str(b).strip()
    if a == b:
        return True
    try:
        return float(a.replace(",", "")) == float(b.replace(",", ""))
    except ValueError:
        return False


# ===== PHAN WINDOWS (chi chay tren Windows) =====
user32 = kernel32 = None
EnumProc = None
WM_GETTEXT = 0x000D
EM_SETSEL = 0x00B1
SMTO_ABORTIFHUNG = 0x0002
KEYEVENTF_KEYUP = 0x0002
KEYEVENTF_UNICODE = 0x0004
INPUT_KEYBOARD = 1
VK_TAB = 0x09


class KEYBDINPUT(ctypes.Structure):
    _fields_ = [("wVk", wintypes.WORD), ("wScan", wintypes.WORD), ("dwFlags", wintypes.DWORD),
                ("time", wintypes.DWORD), ("dwExtraInfo", wintypes.WPARAM)]


class MOUSEINPUT(ctypes.Structure):
    _fields_ = [("dx", wintypes.LONG), ("dy", wintypes.LONG), ("mouseData", wintypes.DWORD),
                ("dwFlags", wintypes.DWORD), ("time", wintypes.DWORD), ("dwExtraInfo", wintypes.WPARAM)]


class _InputUnion(ctypes.Union):
    _fields_ = [("mi", MOUSEINPUT), ("ki", KEYBDINPUT)]


class INPUT(ctypes.Structure):
    _fields_ = [("type", wintypes.DWORD), ("u", _InputUnion)]


def setup_windows():
    global user32, kernel32, EnumProc
    user32 = ctypes.windll.user32
    kernel32 = ctypes.windll.kernel32
    try:
        user32.SetProcessDPIAware()
    except Exception:
        pass
    EnumProc = ctypes.WINFUNCTYPE(wintypes.BOOL, wintypes.HWND, wintypes.LPARAM)

    user32.GetClassNameW.argtypes = [wintypes.HWND, wintypes.LPWSTR, ctypes.c_int]
    user32.GetWindowTextW.argtypes = [wintypes.HWND, wintypes.LPWSTR, ctypes.c_int]
    user32.GetWindowRect.argtypes = [wintypes.HWND, ctypes.POINTER(wintypes.RECT)]
    user32.IsWindowVisible.argtypes = [wintypes.HWND]
    user32.SetForegroundWindow.argtypes = [wintypes.HWND]
    user32.GetWindowThreadProcessId.argtypes = [wintypes.HWND, ctypes.POINTER(wintypes.DWORD)]
    user32.GetWindowThreadProcessId.restype = wintypes.DWORD
    user32.AttachThreadInput.argtypes = [wintypes.DWORD, wintypes.DWORD, wintypes.BOOL]
    user32.SetFocus.argtypes = [wintypes.HWND]
    user32.GetFocus.restype = wintypes.HWND
    user32.SendInput.argtypes = [wintypes.UINT, ctypes.POINTER(INPUT), ctypes.c_int]
    user32.SendMessageW.argtypes = [wintypes.HWND, wintypes.UINT, wintypes.WPARAM, wintypes.LPARAM]
    user32.SendMessageW.restype = wintypes.LPARAM
    user32.SendMessageTimeoutW.argtypes = [
        wintypes.HWND, wintypes.UINT, wintypes.WPARAM, wintypes.LPARAM,
        wintypes.UINT, wintypes.UINT, ctypes.POINTER(ctypes.c_size_t),
    ]
    user32.SendMessageTimeoutW.restype = wintypes.LPARAM
    kernel32.GetCurrentThreadId.restype = wintypes.DWORD


def get_class(hwnd):
    buf = ctypes.create_unicode_buffer(256)
    user32.GetClassNameW(hwnd, buf, 256)
    return buf.value


def get_title(hwnd):
    buf = ctypes.create_unicode_buffer(512)
    user32.GetWindowTextW(hwnd, buf, 512)
    return buf.value


def get_text(hwnd, maxlen=400):
    buf = ctypes.create_unicode_buffer(maxlen)
    result = ctypes.c_size_t()
    ok = user32.SendMessageTimeoutW(hwnd, WM_GETTEXT, maxlen, ctypes.addressof(buf),
                                    SMTO_ABORTIFHUNG, 300, ctypes.byref(result))
    return buf.value if ok else ""


def find_ecus_window():
    found = []

    def cb(hwnd, _):
        if user32.IsWindowVisible(hwnd) and "ECUS5" in get_title(hwnd).upper():
            found.append(hwnd)
        return True

    user32.EnumWindows(EnumProc(cb), 0)
    return found[0] if found else None


def collect_items(main_hwnd):
    handles = []

    def cb(hwnd, _):
        handles.append(hwnd)
        return True

    user32.EnumChildWindows(main_hwnd, EnumProc(cb), 0)
    items = []
    for h in handles:
        r = wintypes.RECT()
        user32.GetWindowRect(h, ctypes.byref(r))
        items.append(Ctl(h, get_class(h), get_text(h), r.left, r.top, r.right, r.bottom,
                         bool(user32.IsWindowVisible(h))))
    return items


def _send_keys(inputs):
    arr = (INPUT * len(inputs))(*inputs)
    user32.SendInput(len(inputs), arr, ctypes.sizeof(INPUT))


def _key(vk=0, scan=0, flags=0):
    ki = KEYBDINPUT(vk, scan, flags, 0, 0)
    inp = INPUT()
    inp.type = INPUT_KEYBOARD
    inp.u.ki = ki
    return inp


def type_text(text):
    for ch in text:
        code = ord(ch)
        if code > 0xFFFF:
            continue  # ky tu ngoai BMP: bo qua
        _send_keys([_key(0, code, KEYEVENTF_UNICODE),
                    _key(0, code, KEYEVENTF_UNICODE | KEYEVENTF_KEYUP)])
        time.sleep(0.01)


def press_tab():
    _send_keys([_key(VK_TAB, 0, 0), _key(VK_TAB, 0, KEYEVENTF_KEYUP)])


def focus_control(main_hwnd, hwnd):
    """Dat focus vao o nhap cua ECUS; tra ve True neu focus dung o do (an toan: sai thi khong go)."""
    user32.SetForegroundWindow(main_hwnd)
    time.sleep(0.1)
    pid = wintypes.DWORD()
    tid_target = user32.GetWindowThreadProcessId(hwnd, ctypes.byref(pid))
    tid_self = kernel32.GetCurrentThreadId()
    user32.AttachThreadInput(tid_self, tid_target, True)
    try:
        user32.SetFocus(hwnd)
        time.sleep(0.1)
        ok = (user32.GetFocus() == hwnd)
    finally:
        user32.AttachThreadInput(tid_self, tid_target, False)
    return ok


def write_field(main_hwnd, hwnd, text):
    if not focus_control(main_hwnd, hwnd):
        return False
    user32.SendMessageW(hwnd, EM_SETSEL, 0, -1)   # chon het noi dung cu, go de len
    time.sleep(0.05)
    type_text(text)
    time.sleep(0.05)
    press_tab()                                   # roi khoi o de ECUS kiem tra / tu dien mo ta
    time.sleep(0.25)
    return True


# ===== CHUONG TRINH CHINH =====
def pick_file():
    files = sorted(OUTPUT_DIR.glob("Template G-*.xlsx"))
    files = [f for f in files if not f.name.startswith("~$")]
    if not files:
        print(f"Khong thay file Template G-*.xlsx trong {OUTPUT_DIR}")
        return None
    print("Cac file Template:")
    for i, f in enumerate(files):
        print(f"  {i}: {f.name}")
    try:
        return files[int(input("\nChon so thu tu file: ").strip())]
    except (ValueError, IndexError):
        print("Lua chon khong hop le.")
        return None


def main():
    if sys.platform != "win32":
        print("Script nay chi chay tren Windows.")
        return
    setup_windows()

    path = pick_file()
    if path is None:
        return

    values = read_template(path)
    print(f"\nDoc file: {path.name}")

    main_hwnd = find_ecus_window()
    if main_hwnd is None:
        print("Khong thay cua so ECUS5 dang mo.")
        return
    items = collect_items(main_hwnd)

    plan = []   # (ten, ctl, gia tri se dien)
    print(f"\n{'Truong':<22}{'Gia tri se dien':<22}{'Hien dang o ECUS':<24}Ghi chu")
    print("-" * 90)
    for name, key, label, idx, fmt in FIELDS:
        raw = values.get(key)
        if raw is None or str(raw).strip() == "":
            print(f"{name:<22}{'(TRONG)':<22}{'':<24}o {CELLS[key]} trong, bo qua "
                  "(mo file Template bang Excel, Ctrl+S neu la cong thuc chua tinh)")
            continue
        try:
            text = fmt(raw)
        except (ValueError, TypeError):
            print(f"{name:<22}{str(raw):<22}{'':<24}khong doi duoc dinh dang, bo qua")
            continue
        ctl, note = locate(items, label, idx)
        if ctl is None:
            print(f"{name:<22}{text:<22}{'':<24}{note}")
            continue
        print(f"{name:<22}{text:<22}{ctl.text.strip()[:22]:<24}{note}")
        plan.append((name, ctl, text))

    if not plan:
        print("\nKhong co truong nao dien duoc.")
        return

    print("\nChua dien gi vao ECUS. Kiem tra bang tren cho dung.")
    if input("Go DIEN de dien that vao ECUS (Enter = thoat): ").strip().upper() != "DIEN":
        print("Da thoat, khong thay doi gi.")
        return

    print("\nBam chuot vao cua so ECUS (tab Thong tin chung) neu can. Bat dau sau 5 giay, "
          "dung cham chuot / ban phim cho den khi xong...")
    time.sleep(5)

    for name, ctl, text in plan:
        ok = write_field(main_hwnd, ctl.hwnd, text)
        if not ok:
            print(f"  [BO QUA] {name}: khong dat duoc focus vao o nay nen khong go (de tranh go nham cho)")
            continue
        now = get_text(ctl.hwnd).strip()
        status = "OK" if same_value(now, text) else f"KHAC (ECUS dang hien {now!r})"
        print(f"  {name:<22} -> {text:<18} {status}")

    print("\nXong. HAY KIEM TRA lai cac o, tu nhap booking / ngay gio, roi moi bam Ghi. "
          "Script khong tu bam Ghi hay Gui.")


if __name__ == "__main__":
    main()