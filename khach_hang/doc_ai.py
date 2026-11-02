"""
Doc hoa don bang AI (Claude API) -> phieu chuan. Dung cho khach moi hoac khach hay doi mau.

Can:
  1. API key cua Anthropic (tao o https://platform.claude.com). Dat vao 1 trong 2 cho:
       - bien moi truong ANTHROPIC_API_KEY, hoac
       - file D:\\test\\khach_hang\\api_key.txt (chi 1 dong la key). KHONG gui file nay cho ai.
  2. Khong can cai them thu vien: goi API bang urllib co san cua Python.

Luu y:
  - Moi lan doc la 1 lan goi API, tinh phi theo luong chu (xem bang gia tren trang Anthropic).
  - Noi dung hoa don duoc gui ra ngoai may. Hoi quy dinh cua cong ty / khach truoc khi dung.
  - AI co the doc sai. Ket qua LUON di qua buoc kiem_tra() (cong tien, nhan don gia) nhu cach quy tac.

Thu rieng: python doc_ai.py "duong\\dan\\hoa_don.docx"
"""
import base64
import json
import os
import urllib.error
import urllib.request
from pathlib import Path

from chung import dong_hang, ngay, phieu_moi

MODEL = "claude-sonnet-5-5"     # doi sang "claude-opus-5-5" neu hoa don kho, "claude-haiku-5-5" de re hon
API_URL = "https://api.anthropic.com/v1/messages"
FILE_KEY = Path(__file__).with_name("api_key.txt")


def _key():
    k = os.environ.get("ANTHROPIC_API_KEY", "").strip()
    if not k and FILE_KEY.exists():
        k = FILE_KEY.read_text(encoding="utf-8").strip()
    if not k:
        raise RuntimeError("Chua co API key: dat bien ANTHROPIC_API_KEY hoac tao file api_key.txt")
    return k


# ===== KHUON KET QUA (JSON schema cua phieu chuan) =====
def _co_the_rong(kieu):
    return {"anyOf": [{"type": kieu}, {"type": "null"}]}


KHUON = {
    "type": "object",
    "properties": {
        "invoice_no": _co_the_rong("string"),
        "invoice_date": {**_co_the_rong("string"), "description": "dang DD/MM/YYYY"},
        "currency": _co_the_rong("string"),
        "incoterm": {**_co_the_rong("string"), "description": "dieu kien + dia diem, vd FOB HO CHI MINH"},
        "total_value": {**_co_the_rong("number"), "description": "tong tri gia hoa don ghi tren hoa don"},
        "packages_qty": _co_the_rong("number"),
        "packages_unit": {**_co_the_rong("string"), "description": "CARTON, BOX, PALLET, BUNDLE, PACKAGE..."},
        "gross_kg": _co_the_rong("number"),
        "consignee": _co_the_rong("string"),
        "buyer": _co_the_rong("string"),
        "pi": _co_the_rong("string"),
        "po": _co_the_rong("string"),
        "payment": {**_co_the_rong("string"), "description": "phuong thuc thanh toan, vd T/T"},
        "origin": {**_co_the_rong("string"), "description": "ma nuoc 2 chu, vd VN"},
        "ship_date": {**_co_the_rong("string"), "description": "dang DD/MM/YYYY"},
        "items": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "ma_hang": _co_the_rong("string"),
                    "mo_ta": _co_the_rong("string"),
                    "ma_hs": _co_the_rong("string"),
                    "so_luong": _co_the_rong("number"),
                    "don_vi": _co_the_rong("string"),
                    "don_gia": {**_co_the_rong("number"), "description": "gia cho 1 don vi"},
                    "tri_gia": _co_the_rong("number"),
                    "xuat_xu": _co_the_rong("string"),
                    "khong_tinh_tien": {"type": "boolean"},
                },
                "required": ["ma_hang", "mo_ta", "ma_hs", "so_luong", "don_vi", "don_gia", "tri_gia",
                             "xuat_xu", "khong_tinh_tien"],
                "additionalProperties": False,
            },
        },
        "ghi_chu_ai": {
            "type": "array", "items": {"type": "string"},
            "description": "nhung cho khong chac chan hoac hoa don ghi khong ro",
        },
    },
    "additionalProperties": False,
}
KHUON["required"] = list(KHUON["properties"])

HUONG_DAN = """Ban doc mot hoa don thuong mai (commercial invoice) de lap to khai hai quan.
Quy tac:
- Chi lay thong tin CO TREN hoa don. Khong co thi de null. Khong doan, khong tu tinh so lieu moi.
- Ngay ghi dang DD/MM/YYYY.
- So ghi dang so (4117.25), khong dau phay ngan cach.
- don_gia la gia cho 1 don vi so_luong. Neu hoa don ghi gia theo 1000 / 2000 cai thi chia ra cho 1 cai.
- Moi dong hang mot phan tu trong items. Dong hang mau / vo thung / "no commercial value" / FOC
  thi khong_tinh_tien = true.
- Khong dua dong tong (Sub Total, Total) vao items.
- Gap cho mo ho, ghi vao ghi_chu_ai."""


# ===== LAY NOI DUNG FILE =====
def _noi_dung(path):
    """Tra ve list cac khoi noi dung de gui API."""
    path = Path(path)
    duoi = path.suffix.lower()
    if duoi in (".docx", ".doc"):
        from doc_word import doc_word
        return [{"type": "text", "text": doc_word(path).van_ban()}]
    if duoi in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        phan = []
        for ws in wb.worksheets:
            phan.append(f"\n[SHEET {ws.title}]")
            for hang in ws.iter_rows(values_only=True):
                o = ["" if v is None else str(v) for v in hang]
                if any(x.strip() for x in o):
                    phan.append(" | ".join(o).rstrip(" |"))
        wb.close()
        return [{"type": "text", "text": "\n".join(phan)}]
    if duoi == ".pdf":
        # Claude doc truc tiep PDF (ca PDF scan)
        data = base64.standard_b64encode(path.read_bytes()).decode("ascii")
        return [{"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": data}}]
    return [{"type": "text", "text": path.read_text(encoding="utf-8", errors="replace")}]


def _goi_api(noi_dung, goi_y, model):
    than = {
        "model": model,
        "max_tokens": 16000,
        "system": HUONG_DAN,
        "messages": [{"role": "user", "content": noi_dung + [
            {"type": "text", "text": "Dac diem rieng cua hoa don khach nay: " + (goi_y or "khong co")}]}],
        "output_config": {"format": {"type": "json_schema", "schema": KHUON}},
    }
    req = urllib.request.Request(
        API_URL, data=json.dumps(than).encode("utf-8"), method="POST",
        headers={"x-api-key": _key(), "anthropic-version": "2023-06-01", "content-type": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=180) as r:
            kq = json.loads(r.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        raise RuntimeError(f"API bao loi {e.code}: {e.read().decode('utf-8', 'replace')[:500]}") from None
    if kq.get("stop_reason") not in ("end_turn", "stop_sequence"):
        raise RuntimeError(f"AI dung giua chung (stop_reason={kq.get('stop_reason')})")
    chu = next(b["text"] for b in kq["content"] if b.get("type") == "text")
    return json.loads(chu)


def chuyen_thanh_phieu(du_lieu):
    """Ket qua JSON cua AI -> phieu chuan (doi ngay sang datetime, gan canh bao)."""
    p = phieu_moi()
    for k in p:
        if k in du_lieu and k not in ("items", "canh_bao"):
            p[k] = du_lieu[k]
    p["invoice_date"] = ngay(du_lieu.get("invoice_date"))
    p["ship_date"] = ngay(du_lieu.get("ship_date"))
    for h in du_lieu.get("items", []):
        p["items"].append(dong_hang(**{k: h.get(k) for k in (
            "ma_hang", "mo_ta", "ma_hs", "so_luong", "don_vi", "don_gia", "tri_gia", "xuat_xu")},
            khong_tinh_tien=bool(h.get("khong_tinh_tien"))))
    p["canh_bao"].append("Doc bang AI: doi chieu ky voi hoa don truoc khi khai")
    p["canh_bao"] += [f"AI ghi chu: {g}" for g in du_lieu.get("ghi_chu_ai", [])]
    return p


def doc_bang_ai(path, goi_y="", model=MODEL):
    return chuyen_thanh_phieu(_goi_api(_noi_dung(path), goi_y, model))


if __name__ == "__main__":
    import sys
    from chung import in_phieu
    if len(sys.argv) < 2:
        print('Cach dung: python doc_ai.py "duong\\dan\\hoa_don"')
        raise SystemExit(1)
    phieu = doc_bang_ai(sys.argv[1])

    class _HoSoTam:
        pass
    in_phieu(phieu, _HoSoTam)
    for c in phieu["canh_bao"]:
        print("  [CANH BAO]", c)
