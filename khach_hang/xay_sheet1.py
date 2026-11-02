"""Dung lai Sheet1 theo BO CUC MOI (02/11/2026) cho 1 file Template. Giu nguyen Sheet2/3/4.
Chay: python xay_sheet1.py <vao.xlsx> <ra.xlsx>

Bo cuc moi:
  A1 = so invoice | A2:G2 trang thai DU / CHUA DU
  3  Phuong thuc van chuyen : B3 Nhu cu/Sua, C3 ma (1-6, 9), D3 ten phuong thuc
  5-8 Mode of Delivery / Forwarder / Ship via / Container Stuffing : C
  10 ma nhan cuoi cung C10 (ten D10) | 11 xep hang C11 (ten D11)
  12 Phuong tien van chuyen : C12 tu dien 9999, D12 ten tau / chuyen
  13 Ky hieu va so hieu : B13
  15 Terms of sale C15 | 16 Currency C16 | 17 Ship to Name C17
  19 So luong kien B19 | 20 Don vi kien B20 (ma D20) | 21 Gross B21
  24 So invoice B24 | 25 Ngay invoice B25
  28 Tri gia gia cong B28 | 29 NVL B29 (=B30-B28) | 30 Tong tri gia hoa don B30
  32 van chuyen bao thue C32 (ten D32)
"""
import sys
from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

VAO, RA = sys.argv[1], sys.argv[2]
DIEN = '"ĐIỀN VÔ"'
def to(m): return PatternFill(start_color=m, end_color=m, fill_type="solid")
VANG, XAM, DO, XANH = to("FFFFFF00"), to("FFDDEBF7"), to("FFFFC7CE"), to("FFC6EFCE")
NHO = Font(italic=True, size=9, color="FF595959")

PT = [("1", "Đường không"), ("2", "Đường biển (container)"), ("3", "Đường biển (không container)"),
      ("4", "Đường bộ (xe tải)"), ("5", "Đường sắt"), ("6", "Đường sông"), ("9", "Loại khác")]
MA = [("C10", "B", "D10", "Nhận cuối cùng"), ("C11", "G", "D11", "Xếp hàng"), ("C32", "L", "D32", "Bảo thuế")]
TU_DONG = [("B19", "Số kiện"), ("D20", "Loại kiện (B20)"), ("B21", "Gross"), ("B24", "Số invoice"),
           ("B25", "Ngày invoice"), ("B30", "Tổng trị giá"), ("C16", "Currency"), ("C17", "Ship to Name")]
SUA = '$B$3="Sửa"'
PT_HOP_LE = "OR(" + ",".join(f'$C$3&""="{m}"' for m, _ in PT) + ")"
BIEN_SONG = 'OR($C$3&""="2",$C$3&""="3",$C$3&""="6")'
CAN_TEN = 'OR($C$3&""="1",$C$3&""="2",$C$3&""="3",$C$3&""="6")'
DK_C3 = f"AND({SUA},NOT({PT_HOP_LE}))"
DK_D12 = f'AND({SUA},{CAN_TEN},$D$12="")'

wb = load_workbook(VAO)
i = wb.sheetnames.index("Sheet1")
wb.remove(wb["Sheet1"])
ws = wb.create_sheet("Sheet1", i)
wb.active = i

nhan = {3: "Phương thức vận chuyển", 5: "Mode of Delivery:", 6: "Forwarder/Consolidator:", 7: "Ship via:",
        8: "Container Stuffing Location:", 10: "mã nhận cuối cùng", 11: "xếp hàng",
        12: "Phương tiện vận chuyển (nhập tay từ booking)", 13: "Ký hiệu và số hiệu", 15: "Terms of sale:",
        16: "Currency:", 17: "Ship to Name:", 19: "Số lượng kiện", 20: "Đơn vị tính kiện hàng",
        21: "Tổng trọng lượng hàng (Gross)", 24: "Số Invoice", 25: "Ngày Invoice", 28: "Tổng trị giá Gia công ",
        29: "Tổng trị giá Nguyên vật liệu", 30: "Tổng trị giá Hóa đơn", 32: "vận chuyển bảo thuế"}
for r, t in nhan.items():
    ws[f"A{r}"] = t
for w, c in ((40, "A"), (12, "B"), (12, "C"), (26, "D"), (22, "E"), (12, "F"), (40, "G")):
    ws.column_dimensions[c].width = w

# --- o nhap (vang) / o tu dong (xanh nhat) ---
for o in ("C3", "C5", "C6", "C7", "C8", "D12", "B13", "C15", "B28"):
    ws[o].fill = VANG
for o in ("B3", "C12"):
    ws[o].fill = XAM
ws["B3"] = "Như cũ"
ws["B3"].font = Font(bold=True)
ws["A1"] = "=B24"
ws["A1"].font = Font(bold=True, size=14)

# --- cong thuc ---
ws["D3"] = ("=IFERROR(CHOOSE(MATCH(C3&\"\",{\"1\",\"2\",\"3\",\"4\",\"5\",\"6\",\"9\"},0),"
            + ",".join(f'"{t}"' for _, t in PT) + '),"ĐIỀN VÔ")')
ws["C12"] = f'=IF(AND({SUA},{BIEN_SONG}),"9999","")'
ws["E12"] = "C12 tự điền: 9999 khi đường biển/sông, trống khi khác | D12: tên tàu / chuyến"
ws["E13"] = "Hồ sơ khách tự ghi (hàm ghi_chu). Để trống = ghép theo mẫu PTTT/TGTT/NNH."
ws["E3"] = "B3: Như cũ = giữ như tờ khai copy lô cũ | Sửa = điền C3 + D12"
for o in ("E3", "E12", "E13"):
    ws[o].font = NHO
for o_ma, cot, o_ten, _ in MA:
    c2 = chr(ord(cot) + 1)
    khac = "OTHER" if o_ma == "C10" else "UNKNOWN"
    ws[o_ten] = (f'=IFERROR(IF(VLOOKUP({o_ma},Sheet4!{cot}:{c2},2,FALSE)="{khac}","ĐIỀN VÔ",'
                 f'VLOOKUP({o_ma},Sheet4!{cot}:{c2},2,FALSE)),"ĐIỀN VÔ")')
ws["C19"], ws["D19"] = "Số lượng", "=B19"
ws["C20"] = "Loại"
ws["D20"] = ('=IF(ISNUMBER(SEARCH("Pallet",B20)),"PP",IF(OR(ISNUMBER(SEARCH("Box",B20)),'
             'ISNUMBER(SEARCH("Boxes",B20))),"BX",IF(OR(ISNUMBER(SEARCH("Bundle",B20)),'
             'ISNUMBER(SEARCH("Package",B20))),"PK","")))')
ws["C21"], ws["D21"] = "KGS", "=B21"
ws["B29"] = "=B30-B28"
ws["B25"].number_format = "dd/mm/yyyy"
for o in ("B28", "B29", "B30"):
    ws[o].number_format = "#,##0.00"

# --- dong trang thai A2 ---
du = ([f'{c}<>""' for c, *_ in MA] + [f"{t}<>{DIEN}" for _, _, t, _ in MA] + [f'{c}<>""' for c, _ in TU_DONG]
      + [f"NOT({DK_C3})", f"NOT({DK_D12})"])
thieu = "".join(f'&IF(OR({c}="",{t}={DIEN}),"{ten} ({c}/{t}); ","")' for c, _, t, ten in MA)
thieu += f'&IF({DK_C3},"Phương thức vận chuyển (C3: 1-6 hoặc 9); ","")'
thieu += f'&IF({DK_D12},"Tên tàu / chuyến (D12); ","")'
thieu += "".join(f'&IF({c}="","{ten}; ","")' for c, ten in TU_DONG)
ws["A2"] = (f'=IF(AND({",".join(du)}),"ĐỦ THÔNG TIN - lưu file (Ctrl+S) rồi chạy ECUS",'
            f'"CHƯA ĐỦ, còn thiếu: "{thieu})')
ws.merge_cells("A2:G2")
ws["A2"].font = Font(bold=True, size=12)
ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[2].height = 48
ws.conditional_formatting.add("A2:G2", FormulaRule(formula=['ISNUMBER(SEARCH("CHƯA ĐỦ",$A$2))'],
                                                   fill=DO, font=Font(bold=True, color="FF9C0006")))
ws.conditional_formatting.add("A2:G2", FormulaRule(formula=['NOT(ISNUMBER(SEARCH("CHƯA ĐỦ",$A$2)))'],
                                                   fill=XANH, font=Font(bold=True, color="FF006100")))
# --- to do o con thieu ---
for c, _, t, _ in MA:
    ws.conditional_formatting.add(f"{c}:{t}", FormulaRule(formula=[f'OR(${c}="",${t}={DIEN})'], fill=DO))
for c, _ in TU_DONG:
    ws.conditional_formatting.add(c, FormulaRule(formula=[f'{c}=""'], fill=DO))
ws.conditional_formatting.add("C3:D3", FormulaRule(formula=[DK_C3], fill=DO))
ws.conditional_formatting.add("D12", FormulaRule(formula=[DK_D12], fill=DO))

# --- chan go sai ---
for c, cot, t, ten in MA:
    dv = DataValidation(type="custom", formula1=f"COUNTIF(Sheet4!${cot}:${cot},{c})>0", allow_blank=True,
                        showErrorMessage=True, errorStyle="stop", errorTitle="Mã không có trong danh mục",
                        error=f"Mã {ten} phải có trong Sheet4 (cột {cot}). Kiểm tra lại mã.",
                        showInputMessage=True, promptTitle=ten,
                        prompt=f"Gõ mã (cột {cot} Sheet4). Mã ZZZ (OTHER/UNKNOWN): gõ tên địa điểm vào {t}.")
    ws.add_data_validation(dv)
    dv.add(c)
for o, ds, tieu_de, loi, goi_y in (
        ("B3", '"Như cũ,Sửa"', "Phương thức + phương tiện vận chuyển", "Chọn Như cũ hoặc Sửa.",
         "Như cũ: giữ đúng như tờ khai copy từ lô cũ, ECUS không đụng vào. Sửa: điền C3 và D12 theo booking "
         "(C12 tự điền), ECUS sẽ điền lại."),
        ("C3", '"1,2,3,4,5,6,9"', "Mã phương thức vận chuyển", "Chỉ nhận 1, 2, 3, 4, 5, 6 hoặc 9.",
         "1 Đường không | 2 Biển container | 3 Biển không container | 4 Bộ | 5 Sắt | 6 Sông | 9 Khác")):
    dv = DataValidation(type="list", formula1=ds, allow_blank=(o != "B3"), showErrorMessage=True, errorStyle="stop",
                        errorTitle=tieu_de, error=loi, showInputMessage=True, promptTitle=tieu_de, prompt=goi_y)
    ws.add_data_validation(dv)
    dv.add(o)
dv = DataValidation(type="custom", formula1="TRUE", allow_blank=True, showInputMessage=True,
                    promptTitle="Tên phương tiện", prompt="Tên tàu / số chuyến theo booking, ví dụ CV7856/25OCT.")
ws.add_data_validation(dv)
dv.add("D12")

wb.save(RA)
print("Da luu", RA)
