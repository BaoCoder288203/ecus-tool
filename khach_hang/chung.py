"""
Phan DUNG CHUNG cho moi khach hang. Thuong khong can sua file nay khi them khach moi.

Y tuong:
    Hoa don khach (Word / Excel / PDF)  --[ho so khach: doc()]-->  PHIEU CHUAN (dict)
    PHIEU CHUAN  --[kiem_tra()]-->  danh sach loi / canh bao
    PHIEU CHUAN  --[ghi_template()]-->  Template <so invoice>.xlsx  --> ecus_fill_demo.py

Moi ho so khach chi can lam 1 viec: doc hoa don cua khach do va dien vao PHIEU CHUAN.
"""
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

# ===== DUONG DAN CHUNG =====
GOC = Path(r"D:\test")                 # thu muc chua code (main.py, ecus_fill_demo.py, khach_hang)
THU_MUC_DON = GOC / "Don hang"         # moi khach + loai to khai 1 thu muc: "Aldila xuat", "Top Glove xuat"...
# Trong moi thu muc khach:  Template_<ten ho so>.xlsx (vd Template_aldila.xlsx)  |  hoa don (co the de trong thu muc con theo lo)  |  output\


def template_goc(ho_so):
    """Template goc cua khach: <thu muc khach>\\Template_<ten ho so>.xlsx (vd Template_aldila.xlsx).
    Con file ten cu 'Template.xlsx' thi tu doi ten sang ten moi."""
    ma = ho_so.__name__.split(".")[-1]
    moi = Path(getattr(ho_so, "TEMPLATE", None) or Path(ho_so.THU_MUC) / f"Template_{ma}.xlsx")
    cu = Path(ho_so.THU_MUC) / "Template.xlsx"
    if not moi.exists() and cu.exists():
        if not la_bo_cuc_moi(cu):
            print(f"[CHU Y] {cu} la bo cuc CU (Phuong thuc van chuyen o dong 12). Khong dung nua: "
                  f"chep file {moi.name} bo cuc moi vao {cu.parent}.")
            return moi
        try:
            cu.rename(moi)
            print(f"Da doi ten {cu.name} -> {moi.name}")
        except OSError as e:
            print(f"[CANH BAO] Khong doi ten duoc {cu} (dang mo trong Excel?): {e}")
            return cu
    return moi


def la_bo_cuc_moi(path):
    """Bo cuc Sheet1 moi (02/11/2026): A3 = 'Phuong thuc van chuyen', A12 = 'Phuong tien van chuyen...'."""
    try:
        wb = load_workbook(path, read_only=True)
        ws = wb["Sheet1"]
        a3 = chuan(ws["A3"].value)
        wb.close()
        return a3 == chuan("Phương thức vận chuyển")
    except Exception:
        return False


def thu_muc_ra(ho_so):
    """Noi luu Template <so invoice>.xlsx cua khach nay (ECUS doc tu day)."""
    return getattr(ho_so, "THU_MUC_RA", None) or Path(ho_so.THU_MUC) / "output"


# ===== PHIEU CHUAN =====
# Day la "ngon ngu chung" giua moi khach hang va phan ECUS. Ho so khach nao cung tra ve dung dang nay.
def phieu_moi():
    return {
        # --- Thong tin chung (dien vao Sheet1 cua Template) ---
        "invoice_no": None,       # so hoa don, vi du "90753570"
        "invoice_date": None,     # ngay hoa don (datetime)
        "currency": None,         # "USD"
        "incoterm": None,         # dieu kien giao hang day du, vi du "FOB HO CHI MINH (CAT LAI), VIETNAM"
        "total_value": None,      # tong tri gia hoa don (so)
        "packages_qty": None,     # so luong kien (so)
        "packages_unit": None,    # loai kien dang chu: "CARTON", "BOX", "PLASTIC PALLET"... (cong thuc Template doi ra ma ECUS)
        "gross_kg": None,         # tong trong luong gross (kg)
        "consignee": None,        # nguoi nhan hang (Ship to Name / CNEE)
        "note": None,             # chuoi 'Ky hieu va so hieu' cho ECUS (ho so khach tu ghep, xem ham ghi_chu)
        # --- Thong tin them, dung de ghep ghi chu / doi chieu (khong bat buoc) ---
        "buyer": None, "pi": None, "po": None, "payment": None,
        "origin": None, "ship_date": None,
        # --- Dong hang ---
        "items": [],              # list cac dict tao bang dong_hang()
        # --- Ghi chu cua bo doc (ho so them vao khi thay dieu bat thuong) ---
        "canh_bao": [],           # dieu can xem lai (muc CANH BAO)
        "loi": [],                # dieu sai ro rang, khong nen khai (muc LOI), vi du INV va PKL khong khop
    }


def so_gon(x, le=8):
    """Bo phan le thua do may tinh (0.009574999999 -> 0.009575) va doi 430000.0 -> 430000."""
    if x is None or isinstance(x, bool):
        return x
    x = round(float(x), le)
    return int(x) if x.is_integer() else x


def dong_hang(ma_hang=None, mo_ta=None, ma_hs=None, so_luong=None, don_vi=None,
              don_gia=None, tri_gia=None, xuat_xu=None, khong_tinh_tien=False):
    """Mot dong hang. don_gia la gia cho 1 don vi (da quy doi neu hoa don ghi gia theo 1000 / 2000 cai)."""
    so_luong, don_gia, tri_gia = so_gon(so_luong), so_gon(don_gia), so_gon(tri_gia, 4)
    return {
        "ma_hang": ma_hang, "mo_ta": mo_ta, "ma_hs": ma_hs, "so_luong": so_luong,
        "don_vi": don_vi, "don_gia": don_gia, "tri_gia": tri_gia, "xuat_xu": xuat_xu,
        "khong_tinh_tien": khong_tinh_tien,   # True = hang mau / vo thung... ghi "no commercial value"
    }


class LoiDoc(Exception):
    """Ho so khach bao loi nay khi hoa don khong dung mau (de chay.py biet ma chuyen sang AI neu duoc phep)."""


# ===== HAM TIEN ICH CHO HO SO KHACH =====
def so(chu, thap_phan="."):
    """Chuyen chu thanh so: '4,117.25' -> 4117.25 ; '' -> None.
    Khach dung dau phay thap phan (vi du '1.234,56') thi goi so(chu, thap_phan=',')."""
    if chu is None:
        return None
    if isinstance(chu, (int, float)):
        return chu
    t = str(chu).strip().replace(" ", "")
    if thap_phan == ",":
        t = t.replace(".", "").replace(",", ".")
    else:
        t = t.replace(",", "")
    t = re.sub(r"[^\d.\-]", "", t)
    if t in ("", "-", ".", "-."):
        return None
    try:
        return float(t)
    except ValueError:
        return None


def ngay(chu, *dinh_dang):
    """Chuyen chu thanh ngay. Mac dinh thu: 21.10.2026, 21/10/2026, 2026-10-21, 21-Oct-26, 21-Oct-2026."""
    if chu is None or isinstance(chu, datetime):
        return chu
    t = str(chu).strip()
    for f in dinh_dang or ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%b-%y", "%d-%b-%Y", "%d %b %Y"):
        try:
            return datetime.strptime(t, f)
        except ValueError:
            pass
    return None


def chuan(s):
    """Chu de so sanh: bo khoang trang thua, bo ':' cuoi, khong phan biet hoa thuong."""
    return re.sub(r"\s+", " ", str(s or "")).strip().rstrip(":").strip().casefold()


# ===== KIEM TRA CHEO =====
BAT_BUOC = ["invoice_no", "invoice_date", "currency", "total_value"]
NEN_CO = ["incoterm", "packages_qty", "packages_unit", "gross_kg", "consignee", "note"]


def hang_se_khai(phieu, ho_so):
    """Cac dong hang se dua vao Template (bo dong khong tinh tien neu ho so khong khai chung)."""
    khai_ca = getattr(ho_so, "KHAI_HANG_KHONG_TINH_TIEN", False)
    return [h for h in phieu["items"] if khai_ca or not h["khong_tinh_tien"]]


def kiem_tra(phieu, ho_so):
    """Tra ve list (muc_do, noi_dung). muc_do = 'LOI' (sai, khong nen khai) hoac 'CANH BAO' (can xem lai)."""
    kq = [("LOI", c) for c in phieu.get("loi", [])]
    kq += [("CANH BAO", c) for c in phieu.get("canh_bao", [])]

    for k in BAT_BUOC:
        if phieu.get(k) in (None, ""):
            kq.append(("LOI", f"Thieu '{k}'"))
    for k in NEN_CO:
        if phieu.get(k) in (None, ""):
            kq.append(("CANH BAO", f"Chua co '{k}' (nhap tay hoac bo sung tu chung tu khac, vi du packing list)"))

    hang = hang_se_khai(phieu, ho_so)
    if not hang:
        kq.append(("LOI", "Khong co dong hang nao"))

    for i, h in enumerate(hang, 1):
        sl, dg, tg = h["so_luong"], h["don_gia"], h["tri_gia"]
        if None in (sl, dg, tg):
            kq.append(("CANH BAO", f"Dong hang {i}: thieu so luong / don gia / tri gia"))
            continue
        lech = abs(sl * dg - tg)
        if lech > max(0.01, abs(tg) * 0.0005):
            kq.append(("LOI", f"Dong hang {i} ({h['mo_ta']}): {sl:,.0f} x {dg} = {sl * dg:,.4f} "
                              f"khac tri gia {tg:,.2f}"))

    tong = sum(h["tri_gia"] or 0 for h in hang)
    if phieu.get("total_value") is not None and hang and abs(tong - phieu["total_value"]) > 0.01:
        kq.append(("LOI", f"Cong cac dong hang = {tong:,.2f} khac tong tri gia hoa don "
                          f"{phieu['total_value']:,.2f}"))
    return kq


# ===== GHI VAO TEMPLATE =====
# Sheet1 cua MOI Template dung chung bo cuc: nhan o cot A, gia tri o cot ghi ben duoi.
# Script tim theo CHU NHAN o cot A, nen ban doi so dong trong Template thoai mai, chi can giu dung chu nhan.
O_SHEET1 = {
    "invoice_no":    ("Số Invoice", "B"),
    "invoice_date":  ("Ngày Invoice", "B"),
    "incoterm":      ("Terms of sale", "C"),
    "currency":      ("Currency", "C"),
    "consignee":     ("Ship to Name", "C"),
    "packages_qty":  ("Số lượng kiện", "B"),
    "packages_unit": ("Đơn vị tính kiện hàng", "B"),
    "gross_kg":      ("Tổng trọng lượng hàng (Gross)", "B"),
    "total_value":   ("Tổng trị giá Hóa đơn", "B"),
    "note":          ("Ký hiệu và số hiệu", "B"),
}
SHEET1 = "Sheet1"


def _sheet(wb, ten):
    for n in wb.sheetnames:
        if n.strip().casefold() == ten.strip().casefold():
            return wb[n]
    return None


def _o_goc(ws, o):
    """O nam trong vung gop -> o goc tren-trai."""
    for r in ws.merged_cells.ranges:
        if o.coordinate in r:
            return ws.cell(r.min_row, r.min_col)
    return o


def ghi_template(phieu, ho_so, ket_qua):
    """Mo Template cua khach, ghi phieu vao, luu thanh file 'ra'. Tra ve list canh bao khi ghi."""
    loi = []
    wb = load_workbook(ho_so.TEMPLATE)

    # --- Sheet1: thong tin chung, tim theo nhan o cot A ---
    ws = _sheet(wb, SHEET1)
    if ws is None:
        raise RuntimeError(f"Template {ho_so.TEMPLATE} khong co sheet {SHEET1}")
    nhan_cot_a = {chuan(ws.cell(r, 1).value): r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value}
    for khoa, (nhan, cot) in O_SHEET1.items():
        gia_tri = phieu.get(khoa)
        if gia_tri in (None, ""):
            continue
        dong = nhan_cot_a.get(chuan(nhan))
        if dong is None:
            loi.append(f"Template thieu dong nhan '{nhan}' o cot A nen chua ghi '{khoa}'")
            continue
        o = _o_goc(ws, ws[f"{cot}{dong}"])
        o.value = gia_tri
        if isinstance(gia_tri, datetime):
            o.number_format = "dd/mm/yyyy"

    # --- Dong hang ---
    ws_h = _sheet(wb, ho_so.SHEET_HANG)
    if ws_h is None:
        loi.append(f"Template khong co sheet '{ho_so.SHEET_HANG}' nen chua ghi dong hang")
    else:
        cot = {k: column_index_from_string(v) for k, v in ho_so.COT_HANG.items()}
        for i, h in enumerate(hang_se_khai(phieu, ho_so)):
            r = ho_so.DONG_BAT_DAU + i
            for k, c in cot.items():
                ws_h.cell(r, c).value = h.get(k)

    # --- Sheet ghi canh bao (de luc mo file van thay) ---
    if phieu.get("_ket_qua_kiem_tra"):
        ws_c = wb.create_sheet("CANH BAO")
        ws_c.append(["Muc do", "Noi dung"])
        for muc, nd in phieu["_ket_qua_kiem_tra"]:
            ws_c.append([muc, nd])
        ws_c.column_dimensions["B"].width = 110

    wb.calculation.fullCalcOnLoad = True
    Path(ket_qua).parent.mkdir(parents=True, exist_ok=True)
    wb.save(ket_qua)
    wb.close()
    return loi


def tinh_lai_bang_excel(cac_file):
    """Mo ngam tung file bang Excel, tinh lai cong thuc, luu lai (can Excel + pywin32).
    Tra ve danh sach file con thieu (dong 2 Sheet1 bao CHUA DU); None neu khong dung duoc Excel."""
    thieu = []
    if not cac_file:
        return thieu
    try:
        import win32com.client as win32
    except Exception:
        print("\n(Chua tinh san cong thuc vi khong dung duoc pywin32. Mo file bang Excel roi Ctrl+S la duoc.)")
        return None
    print("\nDang nho Excel tinh cong thuc (chay ngam)...")
    excel = None
    try:
        excel = win32.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        for f in cac_file:
            wb = None
            try:
                wb = excel.Workbooks.Open(str(Path(f).resolve()))
                excel.CalculateFull()
                wb.Save()
                tt = ""
                try:
                    tt = str(wb.Worksheets("Sheet1").Range("A2").Value or "")
                except Exception:
                    pass
                if "CHƯA ĐỦ" in tt.upper():
                    thieu.append(f)
                print(f"  Da tinh xong: {Path(f).name}" + (f"  -> {tt}" if tt else ""))
            except Exception as e:
                print(f"  LOI khi tinh {Path(f).name}: {e}")
            finally:
                if wb is not None:
                    wb.Close(SaveChanges=False)
    finally:
        if excel is not None:
            excel.Quit()
    return thieu


def in_phieu(phieu, ho_so):
    """In phieu ra man hinh de doi chieu voi hoa don (dung khi viet / thu ho so khach moi)."""
    print("  --- Thong tin chung ---")
    for k in ("invoice_no", "invoice_date", "currency", "incoterm", "total_value", "packages_qty",
              "packages_unit", "gross_kg", "consignee", "buyer", "pi", "po", "payment", "origin",
              "ship_date", "note", "pkl_file", "net_kg", "cbm"):
        v = phieu.get(k)
        if isinstance(v, datetime):
            v = v.strftime("%d/%m/%Y")
        print(f"  {k:<14} {v}")
    print("  --- Dong hang ---")
    for i, h in enumerate(phieu["items"], 1):
        dau = " (khong tinh tien)" if h["khong_tinh_tien"] else ""
        print(f"  {i}. {h['ma_hang']} | {h['mo_ta']} | HS {h['ma_hs']} | {h['so_luong']} {h['don_vi']} "
              f"x {h['don_gia']} = {h['tri_gia']} | {h['xuat_xu']}{dau}")
